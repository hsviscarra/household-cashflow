"""F16: the monthly report (facts, findings by rules or Claude, saved) and the exports."""

import csv
import io
import json
from collections.abc import Iterator
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import Account, Category, MonthlyReport, Transaction
from app.services import export, report

TODAY = "2026-10-28"
USER = {"X-MS-CLIENT-PRINCIPAL-NAME": "Hernan@Example.com"}


@pytest.fixture
def with_key(monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    get_settings.cache_clear()
    yield
    monkeypatch.delenv("ANTHROPIC_API_KEY")
    get_settings.cache_clear()


def _cat(db: Session, name: str) -> int:
    c = db.scalar(select(Category).where(Category.name == name))
    assert c is not None, name
    return c.category_id


def _account(db: Session, name: str) -> int:
    a = db.scalar(select(Account).where(Account.name == name))
    assert a is not None, name
    return a.account_id


def _line(
    db: Session, n: int, account: int, day: date, amount: str, cat: int, kind: str, text: str
) -> None:
    db.add(
        Transaction(
            natural_key=f"rep-{n}",
            account_id=account,
            txn_date=day,
            cash_month=day.replace(day=1),
            amount=Decimal(amount),
            description_raw=text,
            merchant=text,
            type=kind,
            category_id=cat,
            source="manual",
            status="confirmed",
        )
    )


@pytest.fixture
def loaded(client: TestClient, db: Session) -> TestClient:
    """July to October: a salary, groceries, dining that doubles in October, card interest in
    October and a monthly streaming charge from August."""
    bank = _account(db, "Produbanco Ahorros Hernan")
    card = _account(db, "Capital One Venture")
    salary, groceries = _cat(db, "Hernan"), _cat(db, "Supermarket")
    dining, interest = _cat(db, "Restaurants & Cafés"), _cat(db, "Credit Card Interest")
    media = _cat(db, "Media")
    n = 0
    for month in (7, 8, 9, 10):
        n += 1
        _line(db, n, bank, date(2026, month, 15), "3000", salary, "income", "PAYROLL")
        n += 1
        shop = {7: "-400", 8: "-380", 9: "-420", 10: "-410"}[month]
        shop_day = {7: 5, 8: 18, 9: 2, 10: 27}[month]  # irregular: not a subscription
        _line(db, n, card, date(2026, month, shop_day), shop, groceries, "expense", "SUPERMARKET")
        n += 1
        amount = "-300" if month == 10 else "-150"
        _line(db, n, card, date(2026, month, 20), amount, dining, "expense", "BISTRO")
        if month >= 8:
            n += 1
            _line(db, n, card, date(2026, month, 3), "-15.99", media, "expense", "STREAMFLIX")
    n += 1
    _line(db, n, card, date(2026, 10, 25), "-42.10", interest, "fee_interest", "INTEREST CHARGE")
    db.commit()
    return client


def _generate(client: TestClient, month: str = "2026-10") -> Any:
    return client.post(f"/api/reports/{month}/generate", params={"today": TODAY}, headers=USER)


def test_nothing_loaded(client: TestClient) -> None:
    assert client.get("/api/reports").json() == {"months": [], "claude": False}
    assert client.get("/api/reports/2026-10").json() is None
    resp = _generate(client)
    assert resp.status_code == 422 and "Nothing is loaded for October 2026" in resp.json()["detail"]
    assert _generate(client, "2026-11").json()["detail"] == "That month hasn't started yet."
    assert client.get("/api/reports/october").status_code == 422


def test_report_written_by_rules_is_saved(loaded: TestClient, db: Session) -> None:
    assert loaded.get("/api/reports").json()["months"][:2] == ["2026-10-01", "2026-09-01"]
    resp = _generate(loaded)
    assert resp.status_code == 200, resp.text
    r = resp.json()
    assert (r["written_by"], r["model"], r["generated_by"]) == ("rules", None, "hernan@example.com")
    f = r["facts"]
    cf = f["cash_flow"]["month"]
    assert (cf["income"], cf["expenses"], cf["net"]) == ("3000.00", "768.09", "2231.91")
    assert f["cash_flow"]["average_3_months"]["net"] == "2439.34"
    assert f["spending"]["biggest_moves"][0] == {
        "group": "Dining",
        "amount": "300.00",
        "last_month": "150.00",
        "change": "150.00",
    }
    assert f["interest_and_fees"]["interest"] == "42.10"
    assert f["subscriptions"]["active"] == 1 and f["subscriptions"]["monthly_cost"] == "15.99"
    assert [a["group"] for a in f["alerts"]] == ["Dining"]
    assert f["outlook"] is not None and f["outlook"]["from"] == TODAY
    titles = [x["title"] for x in r["findings"]]
    assert 3 <= len(titles) <= 5
    assert titles[0].startswith("October isn't complete yet")  # no statements cover it
    assert "Net +2,231.91 in October" in titles
    # October is incomplete (no statements): a drop or rise against September isn't reported as
    # the month's biggest change; the alert still is.
    assert "Dining above its usual" in titles and not any(t.startswith("Dining up") for t in titles)
    net = next(x for x in r["findings"] if x["title"].startswith("Net"))
    assert "(74% of income)" in net["text"] and net["tone"] == "good"

    # Saved: reading it back doesn't regenerate; regenerating replaces the month's report.
    again = loaded.get("/api/reports/2026-10").json()
    assert again["findings"] == r["findings"] and again["notice"] is None
    assert _generate(loaded).status_code == 200
    assert db.scalar(select(func.count()).select_from(MonthlyReport)) == 1


def test_net_worth_compares_like_with_like(loaded: TestClient) -> None:
    """A car revalued in October changes net worth; an office first valued in October doesn't."""
    car = loaded.post(
        "/api/holdings",
        json={"name": "Car", "holding_class": "vehicle", "value": "10000", "as_of": "2026-09-01"},
    ).json()["holding_id"]
    loaded.post(f"/api/holdings/{car}/values", json={"as_of": "2026-10-10", "amount": "9500"})
    loaded.post(
        "/api/holdings",
        json={
            "name": "Office",
            "holding_class": "property",
            "value": "50000",
            "as_of": "2026-10-20",
        },
    )
    nw = _generate(loaded).json()["facts"]["net_worth"]
    assert nw["change"] == "-500.00" and nw["net_worth"] == "59500.00"
    assert nw["first_valued_this_month"] == [
        {"name": "Office", "side": "asset", "value": "50000.00"}
    ]
    f = report.rule_findings(
        {
            "month": "2026-10",
            "data": {"complete": True, "open": []},
            "cash_flow": {"month": None, "last_month": None, "average_3_months": None},
            "spending": {"biggest_moves": []},
            "alerts": [],
            "outlook": None,
            "interest_and_fees": {
                "interest": "0",
                "fees": "0",
                "interest_year_to_date": "0",
                "fees_year_to_date": "0",
            },
            "subscriptions": {"active": 0, "monthly_cost": "0", "started": [], "stopped": []},
            "net_worth": nw,
        }
    )
    assert f[0].title == "Net worth down 500.00"
    assert f[0].text == (
        "Net worth 59,500.00 on Oct 28, 2026: owned 59,500.00, owed 0.00; first valued this "
        "month: Office (not counted as a change)."
    )


def test_an_older_month_has_no_outlook(loaded: TestClient) -> None:
    r = _generate(loaded, "2026-08").json()
    assert r["facts"]["outlook"] is None
    assert r["facts"]["cash_flow"]["last_month"]["income"] == "3000.00"


def test_report_written_by_claude(
    loaded: TestClient, with_key: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    seen: list[dict[str, Any]] = []

    def fake(f: dict[str, Any]) -> list[report.Finding]:
        seen.append(report.for_claude(f))
        return [report.Finding(f"Finding {i}", "Dining 300.00.", "bad") for i in range(3)]

    monkeypatch.setattr(report, "call_claude", fake)
    assert loaded.get("/api/reports").json()["claude"] is True
    r = _generate(loaded).json()
    assert (r["written_by"], r["model"]) == ("claude", "claude-sonnet-5-5")
    assert [x["title"] for x in r["findings"]] == ["Finding 0", "Finding 1", "Finding 2"]
    # Data Status' texts name accounts (and so people): Claude only gets how many are open.
    assert "open" not in seen[0]["data"] and seen[0]["data"]["open_items"] >= 1
    assert r["facts"]["data"]["open"]  # kept locally

    def broken(f: dict[str, Any]) -> list[report.Finding]:
        raise report.ReportError("Claude could not be reached: APIConnectionError.", 502)

    monkeypatch.setattr(report, "call_claude", broken)
    r = _generate(loaded).json()
    assert r["written_by"] == "rules"
    assert r["notice"] == (
        "Claude could not be reached: APIConnectionError. These findings are written by rules "
        "instead."
    )


def test_call_claude(with_key: None, monkeypatch: pytest.MonkeyPatch) -> None:
    import anthropic

    calls: list[dict[str, Any]] = []
    stop = ["end_turn"]
    three = [{"title": f"T{i}", "text": "x", "tone": "good"} for i in range(3)]
    answer = [json.dumps({"findings": three})]

    class FakeMessages:
        def create(self, **kwargs: Any) -> Any:
            calls.append(kwargs)
            return SimpleNamespace(
                stop_reason=stop[0], content=[SimpleNamespace(type="text", text=answer[0])]
            )

    class FakeClient:
        def __init__(self, **_kw: Any) -> None:
            self.messages = FakeMessages()

    monkeypatch.setattr(anthropic, "Anthropic", FakeClient)
    facts = {"month": "2026-10", "data": {"complete": False, "open": ["Bank X: no statement"]}}
    got = report.call_claude(facts)
    assert [f.title for f in got] == ["T0", "T1", "T2"]
    sent = json.loads(calls[0]["messages"][0]["content"])
    assert sent["data"] == {"complete": False, "open_items": 1}
    assert calls[0]["model"] == "claude-sonnet-5-5"
    assert calls[0]["output_config"]["format"]["type"] == "json_schema"

    stop[0] = "max_tokens"
    with pytest.raises(report.ReportError, match="didn't finish"):
        report.call_claude(facts)
    stop[0] = "end_turn"
    answer[0] = '{"findings": [{"title": "only one", "text": "x", "tone": "good"}]}'
    with pytest.raises(report.ReportError, match="could not be read"):
        report.call_claude(facts)

    class Failing(FakeMessages):
        def create(self, **kwargs: Any) -> Any:
            raise anthropic.APIConnectionError(request=SimpleNamespace())  # type: ignore[arg-type]

    class FailingClient(FakeClient):
        def __init__(self, **_kw: Any) -> None:
            self.messages = Failing()

    monkeypatch.setattr(anthropic, "Anthropic", FailingClient)
    with pytest.raises(report.ReportError, match="could not be reached"):
        report.call_claude(facts)


def test_wording_helpers() -> None:
    assert report._pct(Decimal("-0.08")) == "\u22128%"
    assert report._above(Decimal("0.35")) == "35% above"
    assert report._above(Decimal("2.5")) == "3.5 times"
    assert report._above(Decimal("189.86")) == "191 times"


def test_parse_findings() -> None:
    items = [{"title": f"T{i}", "text": "x", "tone": "odd"} for i in range(7)]
    got = report.parse_findings("Here: " + json.dumps({"findings": items}))
    assert len(got) == 5 and got[0].tone == "neutral"
    for bad in ("nothing", "[1, 2]", '{"x": 1}'):
        with pytest.raises(ValueError):
            report.parse_findings(bad)


def test_rule_findings_cover_the_rest() -> None:
    """Shortfall ahead, subscriptions that started or stopped, net worth down, a loss month."""
    f: dict[str, Any] = {
        "month": "2026-10",
        "data": {"complete": True, "open": []},
        "cash_flow": {
            "month": {
                "income": "100.00",
                "expenses": "300.00",
                "savings": "0.00",
                "debt_payments": "0.00",
                "net": "-200.00",
                "savings_rate": None,
            },
            "last_month": None,
            "average_3_months": None,
        },
        "spending": {
            "biggest_moves": [
                {"group": "Health", "amount": "50.00", "last_month": "80.00", "change": "-30.00"}
            ]
        },
        "alerts": [{"group": "Dining", "amount": "80.00", "average": "0.00", "change": None}],
        "outlook": {"lowest_cash": "-50.00", "lowest_on": "2026-11-14", "next_payday": None},
        "interest_and_fees": {
            "interest": "0.00",
            "fees": "0.00",
            "interest_year_to_date": "0.00",
            "fees_year_to_date": "0.00",
        },
        "subscriptions": {
            "active": 2,
            "monthly_cost": "20.00",
            "started": [{"merchant": "NEWAPP"}],
            "stopped": [{"merchant": "OLDAPP"}],
        },
        "net_worth": {
            "change": "-10.00",
            "net_worth": "5.00",
            "as_of": "2026-10-28",
            "assets": "10.00",
            "liabilities": "5.00",
        },
    }
    got = report.rule_findings(f)
    titles = [x.title for x in got]
    assert titles[0] == "Net \u2212200.00 in October" and got[0].tone == "bad"
    assert titles[1:4] == ["Health down 30.00", "Dining above its usual", "Cash runs short ahead"]
    short = next(x for x in got if x.title == "Cash runs short ahead")
    assert short.text.startswith(
        "The projection reaches \u221250.00 on Nov 14, 2026, and no income is projected after it."
    )
    assert len(got) == 5
    f["spending"]["biggest_moves"] = []
    f["outlook"] = None
    subs = next(x for x in report.rule_findings(f) if x.title == "Subscriptions changed")
    assert subs.text.startswith("Started: NEWAPP; stopped: OLDAPP.")
    f["alerts"], f["outlook"] = [], None
    f["subscriptions"].update(started=[], stopped=[])
    f["net_worth"]["change"] = None
    f["cash_flow"]["month"] = None
    assert [x.title for x in report.rule_findings(f)] == ["2 subscriptions"]


# ---------------------------------------------------------------- exports
def test_export_csv_and_xlsx(loaded: TestClient) -> None:
    keys = [d["key"] for d in loaded.get("/api/exports").json()]
    assert keys == [
        "monthly_cashflow",
        "monthly_cash",
        "category_month",
        "spending_lines",
        "debt_cost",
        "transactions",
    ]
    resp = loaded.get("/api/exports/transactions.csv", params={"month": "2026-10"})
    assert resp.status_code == 200
    assert resp.headers["content-disposition"] == (
        'attachment; filename="cashflow-transactions-2026-10.csv"'
    )
    assert resp.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))
    assert rows[0][:6] == [
        "Transaction id",
        "Txn date",
        "Post date",
        "Cash month",
        "Account",
        "Amount",
    ]
    assert len(rows) == 1 + 5
    interest = next(r for r in rows if r[9] == "INTEREST CHARGE")
    assert (interest[4], interest[5], interest[7]) == (
        "Capital One Venture",
        "-42.10",
        "Financial Costs / Credit Card Interest",
    )

    # A yearly view asked for a month gives the month's year.
    resp = loaded.get("/api/exports/monthly_cashflow.xlsx", params={"month": "2026-10"})
    assert "cashflow-monthly-cashflow-2026.xlsx" in resp.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert ws is not None
    assert [c.value for c in ws[1]][:3] == ["Cash month", "Income", "Expenses"]
    assert ws.freeze_panes == "A2" and ws.max_row == 5
    assert ws["B2"].value == 3000 and ws["B2"].number_format == "#,##0.00;-#,##0.00"
    assert ws["A2"].number_format == "yyyy-mm"
    rate = [c.value for c in ws[1]].index("Savings rate") + 1
    assert ws.cell(2, rate).number_format == "0.0%"

    all_lines = loaded.get("/api/exports/spending_lines.csv").content.decode("utf-8-sig")
    assert len(all_lines.strip().splitlines()) == 1 + 4 * 2 + 3 + 1
    year = loaded.get("/api/exports/debt_cost.csv", params={"year": 2026})
    assert "cashflow-debt-cost-2026.csv" in year.headers["content-disposition"]
    assert loaded.get("/api/exports/nothing.csv").status_code == 404
    assert loaded.get("/api/exports/transactions.pdf").status_code == 422


def test_export_helpers() -> None:
    assert export.filename("transactions", "csv", None, None) == "cashflow-transactions-all.csv"
    t = export.Table(["savings_rate", "flag", "empty"], [(Decimal("0.7473"), True, None)])
    assert export.to_csv(t).decode("utf-8-sig").splitlines()[1] == "0.7473,yes,"
    ws = load_workbook(io.BytesIO(export.to_xlsx(t, "x"))).active
    assert ws is not None and ws["B2"].value == "yes"
    empty = load_workbook(io.BytesIO(export.to_xlsx(export.Table(["a"], []), "x"))).active
    assert empty is not None and empty.max_row == 1
