"""CSV and Excel export of the analysis views (F16, ADR-015).

Each dataset is a mart view (or the transactions with their account and category), optionally
for one month or one year. CSV is UTF-8 with a byte-order mark so Excel opens accents correctly,
amounts as plain numbers with a dot. Excel files keep real numbers (two decimals, thousands
separator) and real dates, with the header row frozen and a filter on it.
"""

import csv
import io
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import Select, select
from sqlalchemy.orm import Session

from app.ingestion.preview import category_paths
from app.models import (
    Account,
    Transaction,
    v_category_month,
    v_debt_cost,
    v_monthly_cash,
    v_monthly_cashflow,
    v_spending_line,
)

FORMATS = ("csv", "xlsx")
MONEY_FORMAT = "#,##0.00;-#,##0.00"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ExportError(Exception):
    def __init__(self, message: str, status: int = 422) -> None:
        super().__init__(message)
        self.status = status


@dataclass
class Dataset:
    key: str
    title: str
    by: str  # month or year: the finest period it can be cut to


DATASETS = {
    d.key: d
    for d in (
        Dataset("monthly_cashflow", "Monthly cash flow (spending view)", "year"),
        Dataset("monthly_cash", "Monthly cash flow (cash view)", "year"),
        Dataset("category_month", "Categories by month", "month"),
        Dataset("spending_lines", "Spending lines", "month"),
        Dataset("debt_cost", "Interest and fees", "month"),
        Dataset("transactions", "Transactions", "month"),
    )
}


@dataclass
class Table:
    columns: list[str]
    rows: Sequence[Sequence[Any]]


def _between(month: date | None, year: int | None) -> tuple[date, date] | None:
    if month is not None:
        return month, month
    if year is not None:
        return date(year, 1, 1), date(year, 12, 1)
    return None


def _view(q: Select[Any], column: Any, period: tuple[date, date] | None, order: Any) -> Select[Any]:
    if period is not None:
        q = q.where(column >= period[0], column <= period[1])
    return q.order_by(*order)


def table(db: Session, key: str, month: date | None = None, year: int | None = None) -> Table:
    d = DATASETS.get(key)
    if d is None:
        raise ExportError(f"There is no export called {key}.", 404)
    if month is not None and d.by == "year":
        year, month = month.year, None  # a yearly view cut to the month's year
    period = _between(month, year)
    if key == "transactions":
        return _transactions(db, period)
    view = {
        "monthly_cashflow": v_monthly_cashflow,
        "monthly_cash": v_monthly_cash,
        "category_month": v_category_month,
        "spending_lines": v_spending_line,
        "debt_cost": v_debt_cost,
    }[key]
    drop = {"group_color", "group_id", "category_id", "account_id", "transaction_id"}
    cols = [c for c in view.c if c.name not in drop]
    c = view.c
    if key == "category_month":
        order = [c.cash_month, c.type, c.group_name, c.category_name]
    elif key == "spending_lines":
        order = [c.txn_date, c.transaction_id]
    elif key == "debt_cost":
        order = [c.cash_month, c.account_name, c.kind]
    else:
        order = [c.cash_month]
    rows = db.execute(_view(select(*cols), view.c.cash_month, period, order)).all()
    return Table([c.name for c in cols], [tuple(r) for r in rows])


def _transactions(db: Session, period: tuple[date, date] | None) -> Table:
    paths = category_paths(db)
    accounts = {a.account_id: a.name for a in db.scalars(select(Account))}
    q = select(Transaction)
    if period is not None:
        q = q.where(Transaction.cash_month >= period[0], Transaction.cash_month <= period[1])
    q = q.order_by(Transaction.txn_date, Transaction.transaction_id)
    columns = [
        "transaction_id",
        "txn_date",
        "post_date",
        "cash_month",
        "account",
        "amount",
        "type",
        "category",
        "merchant",
        "description",
        "person",
        "country",
        "is_fixed",
        "source",
        "status",
        "split_of",
        "notes",
    ]
    rows = [
        (
            t.transaction_id,
            t.txn_date,
            t.post_date,
            t.cash_month,
            accounts.get(t.account_id or -1, "Cash Flow sheet" if t.account_id is None else ""),
            t.amount,
            t.type,
            paths.get(t.category_id or -1),
            t.merchant,
            t.description_raw,
            t.person,
            t.country,
            t.is_fixed,
            t.source,
            t.status,
            t.parent_transaction_id,
            t.notes,
        )
        for t in db.scalars(q)
    ]
    return Table(columns, rows)


def _heading(name: str) -> str:
    return name.replace("_", " ").capitalize()


def _cell(name: str, v: Any) -> Any:
    if isinstance(v, Decimal):
        return format(v if name == "savings_rate" else v.quantize(Decimal("0.01")), "f")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return "yes" if v else "no"
    return "" if v is None else v


def to_csv(t: Table) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow([_heading(c) for c in t.columns])
    w.writerows([_cell(c, v) for c, v in zip(t.columns, r, strict=True)] for r in t.rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(t: Table, title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = title[:31]
    ws.append([_heading(c) for c in t.columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    widths = [len(_heading(c)) for c in t.columns]
    for r in t.rows:
        values = []
        for i, v in enumerate(r):
            if isinstance(v, bool):
                v = "yes" if v else "no"
            values.append(v)
            widths[i] = max(widths[i], min(len(str(v if v is not None else "")), 60))
        ws.append(values)
    for i, name in enumerate(t.columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = widths[i - 1] + 2
        fmt = _format(name, t.rows, i - 1)
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    if t.rows:
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _format(name: str, rows: Iterable[Sequence[Any]], i: int) -> str | None:
    if name == "savings_rate":
        return "0.0%"
    sample = next((r[i] for r in rows if r[i] is not None), None)
    if isinstance(sample, Decimal):
        return MONEY_FORMAT
    if isinstance(sample, date):
        return "yyyy-mm" if name == "cash_month" else "yyyy-mm-dd"
    return None


def filename(key: str, fmt: str, month: date | None, year: int | None) -> str:
    d = DATASETS[key]
    if month is not None and d.by == "month":
        period = month.isoformat()[:7]
    elif month is not None:
        period = str(month.year)
    elif year is not None:
        period = str(year)
    else:
        period = "all"
    return f"cashflow-{key.replace('_', '-')}-{period}.{fmt}"


def render(
    db: Session, key: str, fmt: str, month: date | None = None, year: int | None = None
) -> tuple[bytes, str, str]:
    """(content, media type, file name)."""
    if fmt not in FORMATS:
        raise ExportError("Export as csv or xlsx.")
    t = table(db, key, month, year)
    if fmt == "csv":
        body, media = to_csv(t), "text/csv; charset=utf-8"
    else:
        body, media = to_xlsx(t, DATASETS[key].title), XLSX
    return body, media, filename(key, fmt, month, year)
