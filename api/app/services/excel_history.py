"""Read the Cash Flow workbook into history rows (F2).

The workbook has one sheet per year ("Cash Flow 2021" ... "Cash Flow 2026"), months across row 1,
an INCOME block ending in "Total" and an EXPENSES block ending in "TOTAL". Decisions (2026-10-20):
- Every row value counts; the workbook's TOTAL cells are only compared, because some TOTAL formulas
  leave rows out. Differences are reported with the rows the formula skipped.
- "Cash Flow 2021" holds 29 months, August 2019 to December 2021: the last English month column of
  a sheet is December of the sheet's year, earlier columns count back from it.
- 2026 is actual through June; July to December are a plan and go to the budget table.
- Only the 2026 side tables are read (savings and investments, debts, Johana payments, US transfer
  payments). The 2021 second block and the 2025 Spanish month columns are skipped.
"""

import calendar
import hashlib
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.services.excel_mapping import UNLABELED_ROWS, Mapping, lookup, norm

MONTHS = [m.lower() for m in calendar.month_name[1:]]
SHEET_RE = re.compile(r"^Cash Flow (\d{4})$")
SUM_RE = re.compile(r"^=\s*sum\(\s*([A-Z]+)(\d+)\s*:\s*([A-Z]+)(\d+)\s*\)\s*$", re.IGNORECASE)
ACTUAL_THROUGH = date(2026, 6, 1)
SIDE_TABLE_SHEETS = frozenset({"Cash Flow 2026"})
SIDE_HEADERS = {
    "savings / investments": "savings",
    "deudas": "debts",
    "johana pagos": "empleada_pay",
    "first us transfer payments": "us_transfers",
}
# Labels that end the main table area or start blocks Hernan chose not to import.
SIDE_STOP = {"others", "disposable income mes - mes", "disposable income total"}
CENT = Decimal("0.01")


@dataclass
class Cell:
    sheet: str
    cell: str
    row: int
    block: str  # income, expense, side:<name>
    group: str | None
    label: str
    month: date | None
    value: Decimal
    note: str | None = None


@dataclass
class MonthCheck:
    month: date
    sheet: str
    income_rows: Decimal
    income_workbook: Decimal | None
    expense_rows: Decimal
    expense_workbook: Decimal | None
    left_out: list[str] = field(default_factory=list)
    left_out_total: Decimal = Decimal(0)
    typed_total: bool = False

    @property
    def matches(self) -> bool:
        return (
            self.income_rows == self.income_workbook and self.expense_rows == self.expense_workbook
        )

    @property
    def difference(self) -> Decimal:
        """Row values minus workbook totals (income and expenses combined, absolute)."""
        inc = self.income_rows - (self.income_workbook or Decimal(0))
        exp = self.expense_rows - (self.expense_workbook or Decimal(0))
        return abs(inc) + abs(exp)

    @property
    def explained(self) -> bool:
        """True when the rows the TOTAL formulas left out account for the whole difference."""
        return self.matches or (not self.typed_total and self.difference == self.left_out_total)


@dataclass
class HistoryRow:
    cell: Cell
    mapping: Mapping
    is_plan: bool

    @property
    def signed_amount(self) -> Decimal:
        return self.cell.value if self.mapping.type == "income" else -self.cell.value

    @property
    def natural_key(self) -> str:
        raw = f"excel_history|{self.cell.sheet}|{self.cell.cell}"
        return hashlib.sha256(raw.encode()).hexdigest()


@dataclass
class Snapshot:
    as_of: date
    name: str
    kind: str
    amount: Decimal
    notes: str


@dataclass
class ParsedWorkbook:
    cells: list[Cell]
    rows: list[HistoryRow]
    unmapped: list[Cell]
    checks: list[MonthCheck]
    snapshots: list[Snapshot]

    @property
    def actual_rows(self) -> list[HistoryRow]:
        return [r for r in self.rows if not r.is_plan]

    @property
    def plan_rows(self) -> list[HistoryRow]:
        return [r for r in self.rows if r.is_plan]


def _num(value: Any) -> Decimal | None:
    if isinstance(value, bool) or not isinstance(value, int | float | Decimal):
        return None
    return Decimal(str(value)).quantize(Decimal("0.0001"))


def _month_columns(ws: Worksheet, year: int) -> dict[int, date]:
    """Contiguous English month headers from column B; the last one is December of `year`."""
    cols: list[tuple[int, int]] = []
    col = 2
    while True:
        head = norm(ws.cell(1, col).value)
        if head not in MONTHS:
            break
        cols.append((col, MONTHS.index(head) + 1))
        col += 1
    out: dict[int, date] = {}
    y = year
    for i in range(len(cols) - 1, -1, -1):
        c, m = cols[i]
        if i < len(cols) - 1 and m > cols[i + 1][1]:
            y -= 1
        out[c] = date(y, m, 1)
    return out


def _left_out(formula: Any, rows_with_values: list[int], start: int, end: int) -> list[int] | None:
    """Rows of the block that a =SUM(Xa:Xb) total does not cover; None if not a plain SUM."""
    match = SUM_RE.match(str(formula)) if isinstance(formula, str) else None
    if not match:
        return None
    lo, hi = int(match.group(2)), int(match.group(4))
    return [r for r in rows_with_values if not (lo <= r <= hi) and start < r < end]


def _end_of_month(d: date) -> date:
    return d.replace(day=calendar.monthrange(d.year, d.month)[1])


def parse_workbook(data: bytes) -> ParsedWorkbook:
    values_wb = load_workbook(BytesIO(data), data_only=True, read_only=False)
    formulas_wb = load_workbook(BytesIO(data), data_only=False, read_only=False)
    cells: list[Cell] = []
    checks: list[MonthCheck] = []

    for ws in values_wb.worksheets:
        sheet_match = SHEET_RE.match(ws.title.strip())
        if not sheet_match:
            continue
        fws = formulas_wb[ws.title]
        months = _month_columns(ws, int(sheet_match.group(1)))
        sheet_cells, sheet_checks = _parse_sheet(ws, fws, months)
        cells += sheet_cells
        checks += sheet_checks

    rows: list[HistoryRow] = []
    unmapped: list[Cell] = []
    for c in cells:
        if c.block not in ("income", "expense") or c.month is None or c.value == 0:
            continue
        mapping = lookup(c.block, c.group, c.label)
        if mapping is None:
            unmapped.append(c)
        else:
            rows.append(HistoryRow(c, mapping, is_plan=c.month > ACTUAL_THROUGH))

    return ParsedWorkbook(cells, rows, unmapped, checks, _snapshots(cells))


def _parse_sheet(
    ws: Worksheet, fws: Worksheet, months: dict[int, date]
) -> tuple[list[Cell], list[MonthCheck]]:
    cells: list[Cell] = []
    block: str | None = None
    group: str | None = None
    bounds: dict[str, tuple[int, int]] = {}
    start_row = 0
    after_main = False
    side: str | None = None

    for r in range(2, ws.max_row + 1):
        raw_label = ws.cell(r, 1).value
        label = str(raw_label).strip() if raw_label is not None else ""
        if isinstance(raw_label, datetime):
            label = raw_label.date().isoformat()
        key = norm(label)
        numbers = {c: v for c in months if (v := _num(ws.cell(r, c).value)) is not None}

        if not after_main:
            if key in ("income", "expenses"):
                block, group, start_row = ("income" if key == "income" else "expense"), None, r
                continue
            if block and key in ("total",) and (block == "income" or label == "TOTAL"):
                bounds[block] = (start_row, r)
                if block == "expense":
                    after_main = True
                block = None
                continue
            if block is None:
                continue
            if not numbers:
                group = label or group
                continue
            if not label:
                label = UNLABELED_ROWS.get((ws.title, r), f"(row {r})")
            for c, v in numbers.items():
                cells.append(
                    Cell(ws.title, ws.cell(r, c).coordinate, r, block, group, label, months[c], v)
                )
            continue

        # Side tables after the main blocks (only for the chosen sheets).
        if ws.title not in SIDE_TABLE_SHEETS:
            break
        if key in SIDE_HEADERS:
            side = SIDE_HEADERS[key]
        elif key in SIDE_STOP and side != "us_transfers":
            side = None
            continue
        if side is None or not label or key.startswith("total") or key == "second us transfer":
            continue
        if side == "us_transfers":
            v = _num(ws.cell(r, 2).value)
            if v is not None:
                note = ws.cell(r, 3).value
                cells.append(
                    Cell(
                        ws.title,
                        ws.cell(r, 2).coordinate,
                        r,
                        f"side:{side}",
                        None,
                        label,
                        None,
                        v,
                        str(note) if note else None,
                    )
                )
            continue
        for c, v in numbers.items():
            cells.append(
                Cell(
                    ws.title, ws.cell(r, c).coordinate, r, f"side:{side}", None, label, months[c], v
                )
            )

    return cells, _checks(ws, fws, months, cells, bounds)


def _checks(
    ws: Worksheet,
    fws: Worksheet,
    months: dict[int, date],
    cells: list[Cell],
    bounds: dict[str, tuple[int, int]],
) -> list[MonthCheck]:
    by_col: dict[tuple[str, int], list[Cell]] = defaultdict(list)
    col_of = {m: c for c, m in months.items()}
    for cell in cells:
        if cell.block in ("income", "expense") and cell.month is not None:
            by_col[(cell.block, col_of[cell.month])].append(cell)

    out: list[MonthCheck] = []
    for c, month in months.items():
        if month > ACTUAL_THROUGH:
            continue
        totals: dict[str, tuple[Decimal, Decimal | None]] = {}
        left_out: list[str] = []
        left_total = Decimal(0)
        typed = False
        for block in ("income", "expense"):
            in_block = by_col[(block, c)]
            rows_sum = sum((x.value for x in in_block), Decimal(0))
            if block not in bounds:
                totals[block] = (rows_sum, None)
                continue
            start, end = bounds[block]
            totals[block] = (rows_sum, _num(ws.cell(end, c).value) or Decimal(0))
            formula = fws.cell(end, c).value
            skipped = _left_out(formula, [x.row for x in in_block], start, end)
            if skipped is None:
                is_formula = isinstance(formula, str) and formula.startswith("=")
                typed = typed or (not is_formula and rows_sum != totals[block][1])
            else:
                for x in in_block:
                    if x.row in skipped:
                        left_out.append(f"{x.label} {x.value.quantize(CENT):,}")
                        left_total += abs(x.value)
        if not any(t[0] or t[1] for t in totals.values()):
            continue
        out.append(
            MonthCheck(
                month,
                ws.title,
                totals["income"][0],
                totals["income"][1],
                totals["expense"][0],
                totals["expense"][1],
                left_out,
                left_total,
                typed,
            )
        )
    return out


def _snapshots(cells: list[Cell]) -> list[Snapshot]:
    out: list[Snapshot] = []
    for c in cells:
        if c.block == "side:debts" and c.month is not None:
            out.append(
                Snapshot(
                    _end_of_month(c.month),
                    c.label,
                    "liability",
                    c.value,
                    f"{c.sheet}!{c.cell} (Deudas)",
                )
            )
        elif (
            c.block == "side:savings"
            and c.month is not None
            and norm(c.label) != "savings / investments"
        ):
            out.append(
                Snapshot(
                    _end_of_month(c.month),
                    c.label,
                    "asset",
                    c.value,
                    f"{c.sheet}!{c.cell} (Savings / Investments)",
                )
            )
        elif c.block == "side:us_transfers" and norm(c.label).startswith("remaining to pay"):
            out.append(
                Snapshot(
                    _end_of_month(ACTUAL_THROUGH),
                    c.label,
                    "liability",
                    c.value,
                    f"{c.sheet}!{c.cell} (US transfer payments; date not in the workbook)",
                )
            )
    # The same label can appear twice (two 'Remaining to pay' lines for one lender): number them.
    seen: dict[tuple[date, str], int] = defaultdict(int)
    for s in out:
        seen[(s.as_of, s.name)] += 1
        if seen[(s.as_of, s.name)] > 1:
            s.name = f"{s.name} ({seen[(s.as_of, s.name)]})"
    return out
